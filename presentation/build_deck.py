"""Builds 'Watad Ejust - Academia.pptx' from the Minesweepers 2026 template.

Slide 1 and the last slide are the template's, untouched. Every content slide
reproduces template slide 2 (background, logo, footer) and composes its own
content on top.
"""
import copy
import sys
from pathlib import Path

from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

PRES_DIR = Path("/home/hatem/Development/cpp/minesweeper/presentation")
TEMPLATE = PRES_DIR / "Minesweepers Presentation.pptx"
OUT = PRES_DIR / "Watad Ejust - Academia.pptx"
import tempfile
import zipfile

# Template slide 2's logo, taken straight from the template package.
LOGO = Path(tempfile.mkdtemp()) / "logo.png"
LOGO.write_bytes(zipfile.ZipFile(TEMPLATE).read("ppt/media/image2.png"))

# Palette: template teal + deep night teal, with a sand accent for the desert.
TEAL = RGBColor(0x00, 0x83, 0x8F)
TEAL_LIGHT = RGBColor(0xD6, 0xEC, 0xEE)
CYAN = RGBColor(0x00, 0xAC, 0xC1)
NIGHT = RGBColor(0x0B, 0x2B, 0x30)
INK = RGBColor(0x1E, 0x35, 0x39)
MUTED = RGBColor(0x55, 0x6F, 0x73)
SAND = RGBColor(0xC9, 0x7B, 0x1A)
SAND_LIGHT = RGBColor(0xF6, 0xE6, 0xCF)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BG = RGBColor(0xF0, 0xF7, 0xF7)
LINE = RGBColor(0xC9, 0xDF, 0xE1)
RED = RGBColor(0xC6, 0x28, 0x28)

TITLE_FONT = "Cairo"
BODY_FONT = "Calibri"


def photo_path(stamp):
    return PRES_DIR / f"WhatsApp Image 2026-{stamp}.jpeg"


P1 = photo_path("10-08 at 10.57.42 PM (1)")
P2 = photo_path("10-08 at 10.57.42 PM (2)")
P3 = photo_path("10-08 at 10.57.42 PM")
P4 = photo_path("10-08 at 10.57.43 PM (1)")
P5 = photo_path("10-08 at 10.57.43 PM (2)")
P6 = photo_path("10-08 at 10.57.43 PM (3)")
P7 = photo_path("10-08 at 10.57.43 PM (4)")
P8 = photo_path("10-08 at 10.57.43 PM (5)")
P11 = photo_path("10-08 at 10.57.44 PM (1)")
P12 = photo_path("10-08 at 10.57.44 PM (2)")
P13 = photo_path("10-08 at 10.57.44 PM")
P17 = photo_path("10-10 at 7.38.58 PM")
P19 = photo_path("10-10 at 7.39.29 PM")
P21 = photo_path("10-10 at 7.40.33 PM")
GUI = PRES_DIR / "Screenshot_20261010_211448.png"

prs = Presentation(str(TEMPLATE))
content_layout = prs.slides[1].slide_layout
new_slides = []


# ---------------------------------------------------------------- primitives
def _font(run, size, color, bold=False, font=BODY_FONT, italic=False):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.name = font
    run.font.color.rgb = color


def text(slide, x, y, w, h, content, size=16, color=INK, bold=False,
         font=BODY_FONT, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
         margin=0, spacing=0, name=None):
    """content: str, or list of paragraphs; a paragraph is str or list of
    (text, {size,color,bold,font,italic}) runs."""
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        box.name = name
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.vertical_anchor = anchor
    for side in ("left", "right", "top", "bottom"):
        setattr(tf, f"margin_{side}", Inches(margin))
    paras = content if isinstance(content, list) else [content]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if spacing and i:
            p.space_before = Pt(spacing)
        runs = para if isinstance(para, list) else [(para, {})]
        for t, opts in runs:
            r = p.add_run()
            r.text = t
            _font(r, opts.get("size", size), opts.get("color", color),
                  opts.get("bold", bold), opts.get("font", font),
                  opts.get("italic", False))
    return box


def bullets(slide, x, y, w, h, items, size=16, color=INK, spacing=8):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    for side in ("left", "right", "top", "bottom"):
        setattr(tf, f"margin_{side}", 0)
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        if i:
            p.space_before = Pt(spacing)
        pPr = p._p.get_or_add_pPr()
        pPr.set("marL", str(Inches(0.25)))
        pPr.set("indent", str(-Inches(0.25)))
        bu_clr = etree.SubElement(pPr, qn("a:buClr"))
        etree.SubElement(bu_clr, qn("a:srgbClr")).set("val", "00838F")
        etree.SubElement(pPr, qn("a:buFont")).set("typeface", "Arial")
        etree.SubElement(pPr, qn("a:buChar")).set("char", "•")
        runs = item if isinstance(item, list) else [(item, {})]
        for t, opts in runs:
            r = p.add_run()
            r.text = t
            _font(r, opts.get("size", size), opts.get("color", color),
                  opts.get("bold", False))
    return box


def rect(slide, x, y, w, h, fill=WHITE, line=None, rounded=True, radius=0.08,
         shape=None, line_w=1.0):
    kind = shape or (MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE)
    s = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    if kind == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = min(radius / min(w, h), 0.5)
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(line_w)
    s.shadow.inherit = False
    s.text_frame.text = ""
    return s


def shape_text(s, content, size=14, color=INK, bold=False, align=PP_ALIGN.CENTER,
               anchor=MSO_ANCHOR.MIDDLE, margin=0.08, font=BODY_FONT):
    tf = s.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for side in ("left", "right", "top", "bottom"):
        setattr(tf, f"margin_{side}", Inches(margin))
    paras = content if isinstance(content, list) else [content]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        runs = para if isinstance(para, list) else [(para, {})]
        for t, opts in runs:
            r = p.add_run()
            r.text = t
            _font(r, opts.get("size", size), opts.get("color", color),
                  opts.get("bold", bold), opts.get("font", font))
    return s


def box(slide, x, y, w, h, content, fill=WHITE, line=LINE, size=14, color=INK,
        bold=False, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE):
    s = rect(slide, x, y, w, h, fill=fill, line=line)
    return shape_text(s, content, size, color, bold, align, anchor)


def badge(slide, x, y, label, d=0.42, fill=TEAL, color=WHITE, size=14):
    s = rect(slide, x, y, d, d, fill=fill, shape=MSO_SHAPE.OVAL)
    return shape_text(s, str(label), size, color, True, margin=0)


def arrow(slide, x1, y1, x2, y2, color=TEAL, width=1.75, both=False, dash=False):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1),
                                   Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(width)
    ln = c.line._get_or_add_ln()
    if dash:
        etree.SubElement(ln, qn("a:prstDash")).set("val", "dash")
    if both:
        etree.SubElement(ln, qn("a:headEnd"), type="triangle", w="med", len="med")
    etree.SubElement(ln, qn("a:tailEnd"), type="triangle", w="med", len="med")
    return c


def line(slide, x1, y1, x2, y2, color=LINE, width=0.75):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1),
                                   Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(width)
    return c


def picture(slide, path, x, y, w, h, rounded=True, focus=(0.5, 0.5)):
    """Insert a photo cropped (not stretched) to fill the w x h box."""
    iw, ih = Image.open(path).size
    pic = slide.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(w), Inches(h))
    box_ratio, img_ratio = w / h, iw / ih
    if img_ratio > box_ratio:  # too wide: crop sides
        keep = box_ratio / img_ratio
        left = (1 - keep) * focus[0]
        pic.crop_left, pic.crop_right = left, 1 - keep - left
    else:  # too tall: crop top/bottom
        keep = img_ratio / box_ratio
        top = (1 - keep) * focus[1]
        pic.crop_top, pic.crop_bottom = top, 1 - keep - top
    if rounded:
        pic.auto_shape_type = MSO_SHAPE.ROUNDED_RECTANGLE
        geom = pic._element.spPr.find(qn("a:prstGeom"))
        av = geom.find(qn("a:avLst"))
        if av is None:
            av = etree.SubElement(geom, qn("a:avLst"))
        gd = etree.SubElement(av, qn("a:gd"))
        gd.set("name", "adj")
        gd.set("fmla", f"val {int(50000 * min(0.12 / min(w, h) * 2, 1))}")
    return pic


def notes(slide, t):
    slide.notes_slide.notes_text_frame.text = t


# ---------------------------------------------------------------- slide frames
def base_slide():
    """A copy of template slide 2's frame: background, logo, footer."""
    s = prs.slides.add_slide(content_layout)
    for ph in list(s.placeholders):
        ph._element.getparent().remove(ph._element)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = BG
    s.shapes.add_picture(str(LOGO), Emu(10220800), Emu(233400), Emu(1790225), Emu(1195351))
    text(s, 0.667, 6.73, 6.0, 0.25,
         [[("Minesweepers 13th Edition", {"bold": True}),
           ("  ·  Watad Ejust", {"bold": True})]],
         size=10.5, color=TEAL, font=TITLE_FONT)
    num = len(new_slides) + 2
    text(s, 12.0, 6.73, 0.67, 0.25, str(num), size=10.5, color=TEAL, bold=True,
         font=TITLE_FONT, align=PP_ALIGN.RIGHT)
    new_slides.append(s)
    return s


def content_slide(kicker, title):
    s = base_slide()
    text(s, 0.667, 0.62, 9.0, 0.3, kicker.upper(), size=12, color=SAND, bold=True)
    text(s, 0.667, 0.88, 10.0, 0.75, title, size=32, color=TEAL, bold=True,
         font=TITLE_FONT, anchor=MSO_ANCHOR.TOP)
    return s


def divider(number, title, topics, photo, focus=(0.5, 0.5)):
    s = base_slide()
    picture(s, photo, 7.35, 1.75, 5.3, 4.6, focus=focus)
    text(s, 0.667, 1.55, 3.0, 1.5, number, size=96, color=CYAN, bold=True,
         font=TITLE_FONT)
    text(s, 0.667, 3.0, 6.4, 1.3, title, size=36, color=TEAL, bold=True,
         font=TITLE_FONT, anchor=MSO_ANCHOR.TOP)
    bullets(s, 0.7, 4.45, 6.3, 1.9, topics, size=16, color=MUTED, spacing=6)
    return s


def style_chart(chart, size=12):
    chart.font.size = Pt(size)
    chart.font.name = BODY_FONT
    chart.font.color.rgb = INK


# ================================================================ SLIDES
# 2 — Mission and our answer
s = content_slide("The challenge", "The mission and our answer")
rows = [
    ("Rough 20 × 20 m arena: sand, grass, inclines, ditches",
     "Tracked chassis with 15 cm wide tracks"),
    ("Buried 10 cm cubes (≤ 10 cm deep) and 5 cm surface cubes",
     "Metal detector + camera with on-laptop vision"),
    ("Light and siren on every detection",
     "Light + siren fire on every metal hit"),
    ("Automatic mine map on a 19 × 19 grid (A–S, 1–19)",
     "Kalman localization + automatic mapping"),
    ("Collect the surface mines",
     "Autonomous pickup sequence: arm + electromagnet"),
    ("No Wi-Fi and no power provided at the arena",
     "Onboard access point, 2 × 10 Ah LiPo"),
]
tx, ty, cw = 0.667, 1.85, (6.3, 5.7)
text(s, tx, ty, cw[0], 0.35, "The arena demands", size=14, color=MUTED, bold=True)
text(s, tx + cw[0] + 0.0, ty, cw[1], 0.35, "Our design response", size=14, color=TEAL, bold=True)
rh = 0.66
for i, (req, ans) in enumerate(rows):
    y = ty + 0.42 + i * rh
    rect(s, tx, y, cw[0] - 0.45, rh - 0.1, fill=WHITE, line=LINE)
    text(s, tx + 0.15, y, cw[0] - 0.7, rh - 0.1, req, size=14, color=INK,
         anchor=MSO_ANCHOR.MIDDLE)
    arrow(s, tx + cw[0] - 0.4, y + (rh - 0.1) / 2, tx + cw[0] - 0.05, y + (rh - 0.1) / 2,
          color=SAND, width=2)
    rect(s, tx + cw[0], y, cw[1], rh - 0.1, fill=TEAL_LIGHT)
    text(s, tx + cw[0] + 0.15, y, cw[1] - 0.3, rh - 0.1, ans, size=14, color=NIGHT,
         bold=True, anchor=MSO_ANCHOR.MIDDLE)
notes(s, "Source: Minesweepers 2026 rulebook, Academia category. Each design "
         "choice in this talk traces back to one of these requirements.")

# 3 — Robot at a glance
s = content_slide("Overview", "Robot at a glance")
picture(s, P6, 0.667, 1.85, 6.6, 4.6, focus=(0.5, 0.55))
stats = [
    ("80 × 75 × 46 cm", "Length × width × height"),
    ("2 × 250 W", "24 V geared track motors"),
    ("Pi 4 + STM32", "High-level + real-time control"),
    ("2 × 10 Ah", "LiPo battery packs"),
    ("Own Wi-Fi", "Onboard access point"),
    ("ROS 2", "Robot and base station"),
]
for i, (big, small) in enumerate(stats):
    col, row = i % 2, i // 2
    x, y = 7.6 + col * 2.6, 1.85 + row * 1.55
    rect(s, x, y, 2.45, 1.4, fill=WHITE, line=LINE)
    text(s, x + 0.18, y + 0.2, 2.15, 0.6, big, size=22, color=TEAL, bold=True,
         font=TITLE_FONT)
    text(s, x + 0.18, y + 0.82, 2.15, 0.45, small, size=12, color=MUTED)
notes(s, "Teleoperated tracked UGV. Final mass is still being measured; the rule "
         "limit is 40 kg.")

# 4 — Divider: mechanical
divider("01", "Mechanical Systems & Robot Structure",
        ["Body structure", "Why tracks over wheels on sand",
         "Track construction and drive train", "Arm and autonomous pickup",
         "Built for sand and heat"], P11)

# 5 — Body structure
s = content_slide("Mechanical", "Body structure")
picture(s, P4, 0.667, 1.85, 5.9, 2.75, focus=(0.5, 0.5))
picture(s, P19, 6.767, 1.85, 5.9, 2.75, focus=(0.5, 0.6))
cards = [
    ("Two track pods", "Each pod carries its own motor, chain drive, idlers and road wheels."),
    ("Central spine", "Bridges the pods and carries the electronics, control panel and arm."),
    ("Low and wide", "80 × 75 × 46 cm footprint keeps the centre of gravity low between the tracks."),
]
for i, (h, b) in enumerate(cards):
    x = 0.667 + i * 4.05
    rect(s, x, 4.85, 3.9, 1.6, fill=WHITE, line=LINE)
    badge(s, x + 0.2, 5.05, i + 1)
    text(s, x + 0.75, 5.05, 3.0, 0.42, h, size=16, color=NIGHT, bold=True,
         anchor=MSO_ANCHOR.MIDDLE)
    text(s, x + 0.2, 5.55, 3.55, 0.85, b, size=13, color=MUTED)
notes(s, "Left: rear view with both track loops and the spine. Right: bare chassis "
         "showing idlers and road wheels.")

# 6 — Tracks over wheels
s = content_slide("Mechanical", "Why tracks over wheels on sand")
cd = CategoryChartData()
cd.categories = ["20 kg", "30 kg", "40 kg (rule limit)"]
track_area = 2 * 0.60 * 0.15  # m², two tracks, ~0.6 m ground contact each
wheel_area = 4 * 0.10 * 0.05  # m², four wheels, ~10 × 5 cm patch each
cd.add_series("Tracks", [round(m * 9.81 / track_area / 1000, 1) for m in (20, 30, 40)])
cd.add_series("4 wheels", [round(m * 9.81 / wheel_area / 1000, 1) for m in (20, 30, 40)])
gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.667), Inches(1.85),
                        Inches(6.4), Inches(4.25), cd)
ch = gf.chart
style_chart(ch)
ch.has_title = True
ch.chart_title.text_frame.text = "Ground pressure (kPa) vs robot mass"
_font(ch.chart_title.text_frame.paragraphs[0].runs[0], 14, INK, True)
ch.has_legend = True
ch.legend.position = XL_LEGEND_POSITION.TOP
ch.legend.include_in_layout = False
plot = ch.plots[0]
plot.gap_width = 60
plot.overlap = -10
plot.has_data_labels = True
plot.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
plot.data_labels.font.size = Pt(12)
plot.data_labels.number_format = "0.0"
plot.data_labels.number_format_is_linked = False
for ser, col in zip(plot.series, (TEAL, SAND)):
    ser.format.fill.solid()
    ser.format.fill.fore_color.rgb = col
va = ch.value_axis
va.has_major_gridlines = True
va.major_gridlines.format.line.color.rgb = LINE
va.format.line.fill.background()
va.tick_labels.font.color.rgb = MUTED
ch.category_axis.tick_labels.font.color.rgb = MUTED
ch.category_axis.format.line.color.rgb = LINE
text(s, 0.667, 6.12, 6.4, 0.3,
     "Contact estimates: 2 tracks × 0.60 m × 0.15 m vs 4 wheels × 10 × 5 cm patch.",
     size=10, color=MUTED)
text(s, 7.5, 1.85, 5.2, 0.9,
     [[("≈ 9× ", {"size": 40, "bold": True, "color": TEAL, "font": TITLE_FONT}),
       ("lower ground pressure", {"size": 18, "bold": True, "color": NIGHT})]],
     anchor=MSO_ANCHOR.MIDDLE)
bullets(s, 7.5, 2.9, 5.15, 3.5, [
    [("Floats on loose sand: ", {"bold": True}), ("the long footprint spreads the load instead of digging in.", {})],
    [("Traction: ", {"bold": True}), ("every slat is a grouser biting into sand and grass.", {})],
    [("Bridges ditches and stones ", {"bold": True}), ("that would trap a small wheel.", {})],
    [("Trade-off: ", {"bold": True, "color": SAND}), ("skid-steer turns need more torque, so we use 250 W geared motors.", {})],
], size=15, spacing=10)
notes(s, "Ground pressure p = m·g / contact area. The 0.6 m contact length per "
         "track is an estimate from the 195 cm track loop; replace it with the "
         "measured value. The comparison holds at any final mass.")

# 7 — Track construction
s = content_slide("Mechanical", "Track construction")
steps = [
    (P21, "Aluminium slats", "Extruded profile with a stepped edge that acts as a grouser.", (0.5, 0.5)),
    (P2, "Chain drive", "The chain carries motor torque from the sprocket to every slat.", (0.5, 0.5)),
    (P3, "Rubber grip layer", "Rubber sheet with cuts so it fits around the slats and grips the sand.", (0.5, 1.0)),
]
for i, (img, h, b, f) in enumerate(steps):
    x = 0.667 + i * 4.05
    picture(s, img, x, 1.85, 3.9, 2.85, focus=f)
    badge(s, x + 0.15, 1.98, i + 1, fill=SAND)
    text(s, x, 4.85, 3.9, 0.4, h, size=17, color=NIGHT, bold=True)
    text(s, x, 5.25, 3.9, 0.75, b, size=14, color=MUTED)
    if i < 2:
        arrow(s, x + 3.92, 3.28, x + 4.03, 3.28, color=SAND, width=2.5)
rect(s, 0.667, 6.0, 12.0, 0.5, fill=TEAL_LIGHT)
text(s, 0.667, 6.0, 12.0, 0.5,
     [[("Track loop 195 cm  ·  width 15 cm  ·  ", {"bold": True}),
       ("white idlers on top, road wheels below carry the load", {})]],
     size=14, color=NIGHT, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
notes(s, "Two layers with two jobs: the chain transmits torque, the rubber sheet "
         "adds grip. Cuts in the sheet let it sit flush around each slat.")

# 8 — Drive train
s = content_slide("Mechanical", "Drive train")
picture(s, P17, 0.667, 1.85, 3.6, 4.6, focus=(0.5, 0.45))
tbl_rows = [
    ("Motor", "MY1016Z brushed DC gear motor, 2 units"),
    ("Origin", "Electric wheelchair drive, sourced locally"),
    ("Supply", "24 V DC"),
    ("Rated power", "250 W each, 500 W total"),
    ("Gearbox", "Integrated reduction, high output torque"),
    ("Current", "≈ 10.4 A each at rated output (250 W ÷ 24 V)"),
    ("Transmission", "Chain from gearbox shaft to track sprocket"),
]
gt = s.shapes.add_table(len(tbl_rows), 2, Inches(4.6), Inches(1.85), Inches(8.07),
                        Inches(0.42 * len(tbl_rows))).table
gt.columns[0].width = Inches(2.0)
gt.columns[1].width = Inches(6.07)
for r, (k, v) in enumerate(tbl_rows):
    for c, val in enumerate((k, v)):
        cell = gt.cell(r, c)
        cell.fill.solid()
        cell.fill.fore_color.rgb = TEAL_LIGHT if c == 0 else WHITE
        cell.margin_left = Inches(0.12)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf = cell.text_frame
        tf.text = ""
        run = tf.paragraphs[0].add_run()
        run.text = val
        _font(run, 14, NIGHT if c == 0 else INK, c == 0)
tblPr = gt._tbl.tblPr
tblPr.set("firstRow", "0")
tblPr.set("bandRow", "0")
rect(s, 4.6, 5.15, 8.07, 1.3, fill=WHITE, line=LINE)
text(s, 4.8, 5.25, 7.7, 1.1, [
    [("Skid steering", {"bold": True, "color": TEAL, "size": 16})],
    [("Same speed on both tracks drives straight; a speed difference turns the robot, "
      "down to turning in place. The STM32 mixes the forward and turn commands.", {})],
], size=14, color=INK, spacing=4)
notes(s, "Specs from the MY1016Z listing (24 V, 250 W). Repurposed wheelchair "
         "motors gave high torque on a student budget.")

# 9a — Arm design
s = content_slide("Mechanical · disposal", "Arm design: one joint, free-hanging end")
rect(s, 0.667, 1.85, 6.0, 4.6, fill=WHITE, line=LINE)
gy = 5.95  # ground
line(s, 0.9, gy, 6.45, gy, color=SAND, width=2)
rect(s, 1.0, 4.75, 2.9, 0.9, fill=NIGHT, rounded=True, radius=0.1)
text(s, 1.0, 4.75, 2.9, 0.9, "Chassis", size=13, color=WHITE, bold=True,
     align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
px, py = 3.55, 4.6  # joint
# ghost: arm upright
line(s, px, py, 4.25, 2.45, color=MUTED, width=4)
line(s, 4.25, 2.45, 4.25, 3.15, color=MUTED, width=2)
rect(s, 4.05, 3.15, 0.4, 0.2, fill=MUTED, rounded=False)
text(s, 4.45, 2.2, 1.4, 0.3, "upright", size=11, color=MUTED, bold=True)
# arm lowered
tx, ty = 5.55, 4.15
line(s, px, py, tx, ty, color=TEAL, width=6)
line(s, tx, ty, tx, 5.3, color=NIGHT, width=2.5)
rect(s, tx - 0.3, 5.3, 0.6, 0.22, fill=TEAL, rounded=False)
rect(s, tx - 0.2, 5.55, 0.4, 0.4, fill=SAND, rounded=False)
arrow(s, 4.55, 2.75, 5.35, 3.85, color=SAND, width=2, dash=True)
badge(s, px - 0.21, py - 0.21, "", d=0.42, fill=CYAN)
arrow(s, 6.15, 4.3, 6.15, 5.1, color=NIGHT, width=1.5)
text(s, 6.2, 4.55, 0.35, 0.3, "g", size=14, color=NIGHT, bold=True, font="Cambria")
text(s, 0.9, 3.85, 2.55, 0.65, [[("Single joint", {"bold": True, "color": CYAN})],
                                [("driven by the wiper motor", {"size": 11, "color": MUTED})]],
     size=13, align=PP_ALIGN.RIGHT)
text(s, 4.0, 5.3, 1.15, 0.6, [[("Magnet", {"bold": True, "color": TEAL})],
                              [("+ mine", {"size": 11, "color": SAND})]],
     size=12, align=PP_ALIGN.RIGHT)
text(s, 3.95, 4.58, 1.5, 0.7, [[("Free-hanging end", {"bold": True, "color": NIGHT})],
                               [("gravity keeps the magnet down", {"size": 11, "color": MUTED})]],
     size=12, align=PP_ALIGN.RIGHT)
cards = [
    ("One degree of freedom", "A single rotary joint driven by the wiper motor: one motor, one limit switch, no arm kinematics to solve."),
    ("Free-weight end", "The electromagnet hangs from a free pivot, so gravity keeps it facing the ground at any arm angle, even on a slope."),
    ("Simple, effective control", "Only two positions, lowered and upright, with the upright switch as the reference: easy for the STM32 to control reliably."),
]
for i, (h, b) in enumerate(cards):
    y = 1.85 + i * 1.57
    rect(s, 6.95, y, 5.72, 1.42, fill=TEAL_LIGHT if i == 2 else WHITE, line=LINE)
    badge(s, 7.15, y + 0.18, i + 1, d=0.42, size=14)
    text(s, 7.75, y + 0.15, 4.8, 0.45, h, size=17, color=NIGHT, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    text(s, 7.15, y + 0.66, 5.35, 0.7, b, size=13, color=MUTED)
notes(s, "Design choice: fewer actuators and sensors means fewer failure points in "
         "sand and heat, and a control problem small enough for the STM32 to run "
         "autonomously.")

# 9 — Pickup arm
s = content_slide("Mechanical · disposal", "Autonomous pickup sequence")
seq = [
    ("Trigger", "Operator presses Pick up (J) once. The GUI sends it only if the arm reads upright."),
    ("Lower", "The wiper motor swings the arm down over the surface mine."),
    ("Grab", "The electromagnet is energised and the arm lifts the steel cube."),
    ("Stow", "The magnet releases the mine into the container; the arm returns upright."),
]
for i, (h, b) in enumerate(seq):
    x = 0.667 + i * 3.05
    rect(s, x, 1.85, 2.85, 2.1, fill=WHITE, line=LINE)
    badge(s, x + 0.2, 2.05, i + 1, d=0.5, size=16)
    text(s, x + 0.85, 2.05, 1.9, 0.5, h, size=18, color=NIGHT, bold=True,
         anchor=MSO_ANCHOR.MIDDLE)
    text(s, x + 0.2, 2.7, 2.5, 1.2, b, size=14, color=MUTED)
    if i < 3:
        arrow(s, x + 2.87, 2.9, x + 3.03, 2.9, color=SAND, width=2.5)
rect(s, 0.667, 4.2, 7.0, 2.25, fill=TEAL_LIGHT)
text(s, 0.87, 4.35, 6.6, 2.0, [
    [("Arm motor: WD1160-series wiper motor", {"bold": True, "color": NIGHT, "size": 16})],
    [("12 V DC · aluminium die-cast gearbox · two speeds", {})],
    [("Rated 21.5–28 rpm, 5.2–13.2 N·m (peak up to 15 N·m) depending on variant and speed", {})],
], size=15, color=INK, spacing=6)
rect(s, 7.92, 4.2, 4.75, 2.25, fill=SAND_LIGHT)
text(s, 8.12, 4.35, 4.4, 2.0, [
    [("Runs autonomously", {"bold": True, "color": NIGHT, "size": 16})],
    [("Once triggered, the STM32 runs every step with no operator input. "
      "An upright limit switch blocks a new pickup until the arm is back.", {})],
], size=15, color=INK, spacing=6)
notes(s, "Datasheet: M&C WD1160 / WD1160-B wiper motor, performance at rated 12 V.")

# 10 — Sand and heat
s = content_slide("Mechanical · rough environment", "Built for sand and heat")
picture(s, P7, 0.667, 1.85, 4.0, 2.2, focus=(0.5, 0.55))
picture(s, P1, 0.667, 4.25, 4.0, 2.2, focus=(0.5, 0.5))
for i, (title, items, accent) in enumerate([
    ("Sand and rough terrain", [
        "Low ground pressure keeps the robot on top of loose sand",
        "Slat grousers + rubber grip layer for traction on sand and grass",
        "Long track footprint bridges ditches and stones",
        "Tested outdoors on sand and grass",
    ], TEAL),
    ("Heat", [
        "Electronics sit inside the chassis, shaded from direct sun",
        "Drive motors cruise well below their 250 W rating",
        "LiPo temperature checked between runs; spare pack ready",
    ], SAND),
]):
    x = 4.95 + i * 3.9
    rect(s, x, 1.85, 3.72, 4.6, fill=WHITE, line=LINE)
    text(s, x + 0.25, 2.0, 3.3, 0.5, title, size=18, color=accent, bold=True)
    bullets(s, x + 0.25, 2.6, 3.25, 3.7, items, size=14, spacing=9)
notes(s, "Confirm the heat measures with the team before presenting; the arena is "
         "outdoors in Egypt.")

# 11 — Divider: electronics
divider("02", "Electronic Components & Control System",
        ["Sensors and microcontrollers", "Why an STM32 for real-time control",
         "Power architecture and budget", "On-robot and robot-to-station links",
         "Layered safety stops"], P5, focus=(0.5, 0.5))

# 12 — Sensors and controllers
s = content_slide("Electronics", "Sensors and controllers")
comp = [
    ("Raspberry Pi 4", "ROS 2, localization, video streaming", "Main computer"),
    ("STM32 Black Pill", "Motor PWM, encoders, ADC, arm, magnet, watchdog", "UART to Pi"),
    ("Wheel encoders", "Track odometry", "STM32 timers"),
    ("MPU-9250 IMU", "9-axis: accelerometer, gyro, magnetometer", "I2C to Pi"),
    ("Metal detector", "Buried and surface mines", "Analog to STM32 ADC"),
    ("USB camera", "Surface-mine vision and operator view", "USB to Pi"),
    ("Motor drivers", "2 × 250 W track motors + arm motor", "PWM from STM32"),
    ("Electromagnet + limit switch", "Mine pickup, arm-upright interlock", "STM32 I/O"),
    ("Light + siren", "Signal every detection", "Digital output"),
    ("Wi-Fi access point", "Robot's own network to the base station", "Network"),
]
t = s.shapes.add_table(len(comp) + 1, 3, Inches(0.667), Inches(1.8), Inches(12.0),
                       Inches(0.42 * (len(comp) + 1))).table
for c, wdt in enumerate((3.4, 5.6, 3.0)):
    t.columns[c].width = Inches(wdt)
for r, row in enumerate([("Component", "Role", "Interface")] + comp):
    for c, val in enumerate(row):
        cell = t.cell(r, c)
        cell.fill.solid()
        cell.fill.fore_color.rgb = TEAL if r == 0 else (WHITE if r % 2 else TEAL_LIGHT)
        cell.margin_left = Inches(0.12)
        cell.margin_top = cell.margin_bottom = Inches(0.03)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf = cell.text_frame
        tf.text = ""
        run = tf.paragraphs[0].add_run()
        run.text = val
        _font(run, 13.5, WHITE if r == 0 else (NIGHT if c == 0 else INK), r == 0 or c == 0)
    t.rows[r].height = Inches(0.42)
t._tbl.tblPr.set("bandRow", "0")
notes(s, "Real-time I/O lives on the STM32; anything that needs Linux, ROS or a "
         "camera lives on the Pi.")

# 13 — Why STM32
s = content_slide("Electronics", "Why an STM32 for real-time control")
cards = [
    ("Hardware encoder timers", "Quadrature counting in hardware: no lost counts at full track speed."),
    ("Deterministic timing", "A bare-metal control loop with fixed timing for PWM, sampling and the watchdog."),
    ("12-bit ADC", "Samples the metal-detector signal for clean threshold detection."),
    ("Dedicated to control", "Networking lives on the Pi and the onboard AP; the MCU only runs real-time I/O."),
]
for i, (h, b) in enumerate(cards):
    col, row = i % 2, i // 2
    x, y = 0.667 + col * 6.1, 1.85 + row * 2.35
    rect(s, x, y, 5.9, 2.15, fill=WHITE if i % 3 else TEAL_LIGHT, line=LINE)
    badge(s, x + 0.25, y + 0.25, i + 1, d=0.5, size=16)
    text(s, x + 0.95, y + 0.25, 4.7, 0.5, h, size=18, color=NIGHT, bold=True,
         anchor=MSO_ANCHOR.MIDDLE)
    text(s, x + 0.25, y + 0.95, 5.4, 1.1, b, size=15, color=MUTED)
notes(s, "STM32 Black Pill. It talks to the Pi over UART with COBS framing and "
         "enforces the 250 ms motor watchdog.")

# 14 — Electrical architecture
s = content_slide("Electronics", "Power architecture")
box(s, 0.667, 3.35, 2.1, 1.2, [[("2 × 10 Ah", {"bold": True, "size": 18, "color": NIGHT})],
                               "LiPo packs"], fill=SAND_LIGHT, line=None)
arrow(s, 2.8, 3.95, 3.25, 3.95, color=NIGHT, width=2.5)
box(s, 3.28, 3.35, 2.1, 1.2, [[("Main switch", {"bold": True})], "+ E-stop buttons"],
    fill=RGBColor(0xFB, 0xE3, 0xE3), line=None, color=RED)
arrow(s, 5.41, 3.95, 5.86, 3.95, color=NIGHT, width=2.5)
box(s, 5.89, 3.35, 1.5, 1.2, [[("24 V", {"bold": True, "size": 20, "color": NIGHT})], "bus"],
    fill=TEAL_LIGHT, line=None)
branches = [
    ("Motor drivers", "2 × MY1016Z, 250 W each", 1.85),
    ("24 → 12 V converter", "Wiper arm motor · electromagnet · light + siren", 3.05),
    ("24 → 5 V converter", "Raspberry Pi 4 → camera, IMU, STM32", 4.25),
    ("Access point supply", "Robot Wi-Fi network", 5.45),
]
for h, b, y in branches:
    arrow(s, 7.42, 3.95, 8.05, y + 0.45, color=TEAL, width=1.75)
    box(s, 8.1, y, 4.57, 0.9, [[(h, {"bold": True, "color": NIGHT})], [(b, {"size": 13, "color": MUTED})]],
        fill=WHITE, line=LINE, size=15, align=PP_ALIGN.LEFT)
text(s, 0.667, 5.0, 6.7, 1.4, [
    [("E-stop first: ", {"bold": True, "color": RED}),
     ("every load sits downstream of the hardware stop.", {})],
    [("High-current drive path ", {"bold": True}),
     ("is kept separate from the regulated logic supplies.", {})],
], size=14, color=INK, spacing=6)
notes(s, "Check the pack wiring (series vs parallel) and converter ratings with the "
         "electronics lead before presenting.")

# 15 — On-robot communication
s = content_slide("Electronics · communication", "Pi to STM32 link")
box(s, 0.667, 1.9, 3.0, 1.3, [[("Raspberry Pi 4", {"bold": True, "size": 18})],
                              [("ROS 2 hardware-interface node", {"size": 13, "color": MUTED})]],
    fill=WHITE, line=LINE)
box(s, 9.67, 1.9, 3.0, 1.3, [[("STM32 Black Pill", {"bold": True, "size": 18})],
                             [("Real-time firmware", {"size": 13, "color": MUTED})]],
    fill=WHITE, line=LINE)
arrow(s, 3.75, 2.3, 9.6, 2.3, color=TEAL, width=2.5)
arrow(s, 9.6, 2.85, 3.75, 2.85, color=SAND, width=2.5)
text(s, 3.9, 1.88, 5.6, 0.35, "Down: drive (forward, turn) · pickup", size=13, color=TEAL,
     bold=True, align=PP_ALIGN.CENTER)
text(s, 3.9, 2.9, 5.6, 0.35, "Up: encoder counts · metal level · arm upright", size=13,
     color=SAND, bold=True, align=PP_ALIGN.CENTER)
text(s, 0.667, 3.6, 8.0, 0.4, "UART frame", size=16, color=NIGHT, bold=True)
fields = [("COBS byte", 1.3, TEAL_LIGHT), ("type", 1.1, TEAL_LIGHT),
          ("payload (packed struct)", 4.6, WHITE), ("checksum", 1.6, SAND_LIGHT),
          ("0x00", 1.1, NIGHT)]
x = 0.667
for label, wdt, fill in fields:
    box(s, x, 4.05, wdt, 0.8, label, fill=fill, line=LINE, size=14,
        color=WHITE if fill == NIGHT else INK, bold=True)
    x += wdt + 0.06
cards = [
    ("COBS framing", "Zero bytes appear only as the delimiter, so the receiver always resynchronises on the next frame."),
    ("Checksum", "A corrupted frame is dropped, never acted on."),
    ("250 ms watchdog", "No valid command for 250 ms → the STM32 sets motor PWM to zero."),
]
for i, (h, b) in enumerate(cards):
    xx = 0.667 + i * 4.05
    rect(s, xx, 5.1, 3.9, 1.35, fill=WHITE, line=LINE)
    text(s, xx + 0.2, 5.18, 3.5, 0.4, h, size=15, color=TEAL, bold=True)
    text(s, xx + 0.2, 5.55, 3.55, 0.85, b, size=13, color=MUTED)
notes(s, "Structs are packed so the 32-bit MCU and the 64-bit Pi agree on the "
         "byte layout.")

# 16 — Robot to station
s = content_slide("Electronics · communication", "Robot to control station: our own network")
rect(s, 0.667, 1.85, 4.2, 4.6, fill=WHITE, line=LINE)
text(s, 0.867, 1.95, 3.8, 0.4, "Robot", size=17, color=TEAL, bold=True)
box(s, 0.967, 2.5, 3.6, 0.9, [[("Onboard Wi-Fi access point", {"bold": True})]],
    fill=TEAL_LIGHT, line=None, size=15, color=NIGHT)
box(s, 0.967, 3.65, 3.6, 0.9, [[("Raspberry Pi 4", {"bold": True})],
                               [("ROS 2 nodes · RTSP server", {"size": 13, "color": MUTED})]],
    fill=BG, line=LINE, size=15)
box(s, 0.967, 4.8, 3.6, 0.9, [[("STM32 + sensors", {"bold": True})]], fill=BG, line=LINE, size=15)
rect(s, 8.47, 1.85, 4.2, 4.6, fill=WHITE, line=LINE)
text(s, 8.67, 1.95, 3.8, 0.4, "Base station (outside the fence)", size=17, color=TEAL, bold=True)
box(s, 8.77, 2.5, 3.6, 0.9, [[("Laptop joins the robot's AP", {"bold": True})]],
    fill=TEAL_LIGHT, line=None, size=15, color=NIGHT)
box(s, 8.77, 3.65, 3.6, 0.9, [[("ROS 2 bridge", {"bold": True})], [("Docker", {"size": 13, "color": MUTED})]],
    fill=BG, line=LINE, size=15)
box(s, 8.77, 4.8, 3.6, 0.9, [[("GUI + vision", {"bold": True})]], fill=BG, line=LINE, size=15)
arrow(s, 4.95, 2.95, 8.4, 2.95, color=TEAL, width=3, both=True)
text(s, 5.05, 2.25, 3.25, 0.65, [[("ROS 2 DDS", {"bold": True, "color": TEAL})],
                                 [("drive at 20 Hz · odometry · metal · arm", {"size": 12, "color": MUTED})]],
     size=14, align=PP_ALIGN.CENTER)
arrow(s, 4.95, 4.15, 8.4, 4.15, color=SAND, width=3)
text(s, 5.05, 4.25, 3.25, 0.75, [[("RTSP video, H.264", {"bold": True, "color": SAND})],
                                 [("kept out of DDS to save bandwidth", {"size": 12, "color": MUTED})]],
     size=14, align=PP_ALIGN.CENTER)
text(s, 5.0, 5.35, 3.4, 1.1, [[("Why on the robot? ", {"bold": True, "color": NIGHT}),
                               ("The arena provides no Wi-Fi, so the link is the same "
                                "in the lab and at the competition.", {})]],
     size=13, color=INK, align=PP_ALIGN.CENTER)
notes(s, "The 20 × 20 m arena diagonal is about 28 m plus the distance to the base "
         "station; the AP rides on the robot so range is set by our own hardware.")

# 17 — Power budget
s = content_slide("Electronics", "Power budget")
loads = [("Track motors (2 × 250 W)", 500), ("Arm wiper motor", 67), ("Electromagnet*", 12),
         ("Raspberry Pi 4", 7.5), ("Wi-Fi access point*", 6), ("Light + siren*", 5),
         ("STM32 + sensors + camera*", 3.5)]
cd = CategoryChartData()
cd.categories = [n for n, _ in reversed(loads)]
cd.add_series("Rated W", [v for _, v in reversed(loads)])
gf = s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(0.667), Inches(1.8),
                        Inches(7.3), Inches(4.35), cd)
ch = gf.chart
style_chart(ch)
ch.has_legend = False
ch.has_title = True
ch.chart_title.text_frame.text = "Rated power per load (W)"
_font(ch.chart_title.text_frame.paragraphs[0].runs[0], 14, INK, True)
plot = ch.plots[0]
plot.gap_width = 45
plot.has_data_labels = True
plot.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
plot.data_labels.font.size = Pt(12)
plot.data_labels.number_format = "0.#"
plot.data_labels.number_format_is_linked = False
ser = plot.series[0]
ser.format.fill.solid()
ser.format.fill.fore_color.rgb = TEAL
pt = ser.points[len(loads) - 1]
pt.format.fill.solid()
pt.format.fill.fore_color.rgb = SAND
va = ch.value_axis
va.has_major_gridlines = True
va.major_gridlines.format.line.color.rgb = LINE
va.format.line.fill.background()
va.tick_labels.font.color.rgb = MUTED
ch.category_axis.tick_labels.font.color.rgb = INK
ch.category_axis.format.line.color.rgb = LINE
text(s, 0.667, 6.15, 7.3, 0.3,
     "Motor values from datasheets (wiper: 12 V × 5.6 A). *Estimates, to be measured.",
     size=10, color=MUTED)
stats = [("≈ 600 W", "Peak, all loads at rating"),
         ("≈ 230 W", "Match average (drive at ~40 %)"),
         ("≈ 75 Wh", "Energy for one 20 min run"),
         ("≈ 3 runs", "From ≈ 222 Wh on board**")]
for i, (big, small) in enumerate(stats):
    col, row = i % 2, i // 2
    x, y = 8.3 + col * 2.2, 1.85 + row * 1.75
    rect(s, x, y, 2.07, 1.6, fill=WHITE if i < 3 else TEAL_LIGHT, line=LINE)
    text(s, x + 0.15, y + 0.2, 1.8, 0.6, big, size=24, color=TEAL if i < 3 else NIGHT,
         bold=True, font=TITLE_FONT)
    text(s, x + 0.15, y + 0.85, 1.85, 0.65, small, size=12, color=MUTED)
text(s, 8.3, 5.4, 4.37, 0.9,
     "**Assumes the two 10 Ah packs form one 22.2 V (6S) bus. "
     "Peak ≈ 27 A, well inside the packs' discharge rating.",
     size=11, color=MUTED)
notes(s, "Replace the starred estimates with measured currents. Battery energy "
         "depends on the pack configuration.")

# 18 — Safety chain
s = content_slide("Control system", "Stopping is layered")
layers = [
    ("GUI focus lost", "Drive zeroed, speed lock cleared"),
    ("GUI refresh", "No UI update for 200 ms → zero"),
    ("ROS bridge", "Drops commands > 200 ms old, publishes one zero"),
    ("Raspberry Pi", "Own motion timeout"),
    ("STM32 watchdog", "250 ms without valid frame → PWM 0"),
    ("Hardware E-stop", "Cuts power to the loads"),
]
for i, (h, b) in enumerate(layers):
    x = 0.667 + i * 2.03
    hw = i == len(layers) - 1
    rect(s, x, 2.3, 1.85, 2.6, fill=RGBColor(0xFB, 0xE3, 0xE3) if hw else WHITE,
         line=None if hw else LINE)
    badge(s, x + 0.7, 2.45, i + 1, fill=RED if hw else TEAL)
    text(s, x + 0.12, 3.0, 1.62, 0.6, h, size=15, color=RED if hw else NIGHT, bold=True,
         align=PP_ALIGN.CENTER)
    text(s, x + 0.12, 3.65, 1.62, 1.2, b, size=13, color=MUTED, align=PP_ALIGN.CENTER)
    if i < len(layers) - 1:
        arrow(s, x + 1.86, 3.6, x + 2.02, 3.6, color=SAND, width=2.5)
text(s, 0.667, 1.82, 12, 0.4, "Laptop  →  robot", size=14, color=MUTED, bold=True)
rect(s, 0.667, 5.25, 12.0, 1.15, fill=TEAL_LIGHT)
text(s, 0.9, 5.25, 11.5, 1.15, [
    [("Every stage fails safe on its own. ", {"bold": True, "color": NIGHT}),
     ("A lost link, hung GUI or crashed process stops the robot, and "
      "reconnecting never resumes motion: the operator must press a key again.", {})]],
     size=15, color=INK, anchor=MSO_ANCHOR.MIDDLE)
notes(s, "These are per-stage timeouts, each independent of the others.")

# 19 — Divider: software
divider("03", "Software Stack",
        ["Architecture: STM32, Pi and laptop", "Mapping and localization",
         "Mine detection and classification", "Control station GUI",
         "Navigation plan and reliability"],
        P13)

# 20 — Architecture
s = content_slide("Software", "Architecture: three tiers, one contract")
cols = [("STM32 firmware", 0.667, 2.7), ("Raspberry Pi 4 · ROS 2 in Docker", 3.65, 4.0),
        ("Base-station laptop", 7.95, 4.72)]
for name, x, w in cols:
    rect(s, x, 1.8, w, 4.65, fill=WHITE, line=LINE)
    text(s, x + 0.15, 1.88, w - 0.3, 0.4, name, size=15, color=TEAL, bold=True)
box(s, 0.867, 2.45, 2.3, 2.0, [[("Real-time loops", {"bold": True})],
                               [("PWM · encoders · ADC · arm sequence · watchdog", {"size": 12, "color": MUTED})]],
    fill=BG, line=None, size=14)
pi_nodes = [("Hardware interface", 2.45), ("IMU filter (Madgwick)", 3.3),
            ("Kalman filter", 4.15), ("RTSP video server", 5.4)]
for label, y in pi_nodes:
    box(s, 3.85, y, 3.1, 0.65, label, fill=BG, line=None, size=14, bold=True)
lap_nodes = [("ROS 2 bridge (Docker)", 2.45), ("Control panel GUI", 3.75), ("Vision: ONNX Runtime", 5.4)]
for label, y in lap_nodes:
    box(s, 8.15, y, 2.6, 0.75, label, fill=BG, line=None, size=14, bold=True)
arrow(s, 3.2, 2.78, 3.82, 2.78, color=NIGHT, both=True)
text(s, 3.1, 2.45, 0.8, 0.28, "UART", size=10, color=NIGHT, bold=True, align=PP_ALIGN.CENTER)
arrow(s, 5.4, 3.13, 5.4, 3.27, color=TEAL)
arrow(s, 5.4, 3.98, 5.4, 4.12, color=TEAL)
arrow(s, 6.98, 4.47, 8.12, 2.82, color=TEAL, both=True)
text(s, 7.0, 2.45, 0.65, 0.5, [[("DDS", {"bold": True})], "Wi-Fi"], size=10, color=TEAL, align=PP_ALIGN.CENTER)
arrow(s, 6.98, 5.72, 8.12, 5.77, color=SAND)
text(s, 7.0, 5.3, 0.65, 0.3, "RTSP", size=10, color=SAND, bold=True, align=PP_ALIGN.CENTER)
arrow(s, 9.45, 3.23, 9.45, 3.72, color=CYAN, both=True)
arrow(s, 9.45, 5.37, 9.45, 4.53, color=CYAN)
text(s, 10.85, 2.4, 1.75, 3.9, [
    [("ZeroMQ IPC", {"bold": True, "color": CYAN})],
    [("One packed, versioned C++ header shared by GUI, bridge and vision", {"size": 12, "color": MUTED})],
    [("", {})],
    [("Separate processes", {"bold": True, "color": NIGHT})],
    [("An inference stall or bridge crash can't freeze the GUI", {"size": 12, "color": MUTED})],
], size=13, spacing=2)
notes(s, "Kalman output is /odometry/filtered. The GUI owns mine mapping; the bridge "
         "only translates.")

# 21 — Mapping and localization
s = content_slide("Software", "Mapping and localization")
gx, gy, gs = 0.95, 2.05, 4.1  # grid origin and size (inches)
n = 19
cell = gs / n
rect(s, gx, gy, gs, gs, fill=WHITE, line=TEAL, rounded=False)
for i in range(1, n):
    line(s, gx + i * cell, gy, gx + i * cell, gy + gs)
    line(s, gx, gy + i * cell, gx + gs, gy + i * cell)
for i, letter in enumerate("ABCDEFGHIJKLMNOPQRS"):
    if i % 3 == 0 or i == n - 1:
        text(s, gx + i * cell - 0.1, gy - 0.28, cell + 0.2, 0.25, letter, size=10,
             color=MUTED, align=PP_ALIGN.CENTER)
for i in range(n):
    if i % 3 == 0 or i == n - 1:
        text(s, gx - 0.38, gy + gs - (i + 1) * cell, 0.32, cell, str(i + 1), size=10,
             color=MUTED, align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)
# sweep path
path_y = [1.5, 4.5, 7.5]
pts = []
for k, row in enumerate(path_y):
    yy = gy + gs - row * cell
    xs = (gx + 0.5 * cell, gx + gs - 0.5 * cell)
    pts.append((xs[k % 2], yy))
    pts.append((xs[(k + 1) % 2], yy))
for (x1, y1), (x2, y2) in zip(pts, pts[1:]):
    arrow(s, x1, y1, x2, y2, color=CYAN, width=1.5, dash=True)
for cx, cy, colr in [(5, 4, SAND), (13, 7, TEAL), (9, 2, SAND)]:
    m = rect(s, gx + (cx - 0.5) * cell, gy + gs - (cy + 0.5) * cell, cell * 0.9, cell * 0.9,
             fill=colr, rounded=False)
badge(s, gx - 0.12, gy + gs - 0.3, "R", d=0.36, fill=NIGHT, size=11)
leg = [("Buried", SAND), ("Surface", TEAL)]
for i, (lab, colr) in enumerate(leg):
    rect(s, gx + i * 1.4, gy + gs + 0.2, 0.2, 0.2, fill=colr, rounded=False)
    text(s, gx + i * 1.4 + 0.28, gy + gs + 0.14, 1.1, 0.32, lab, size=12, color=MUTED)
steps = [
    ("Localize", "Kalman filter fuses track odometry (STM32 encoders) with the MPU-9250 heading. No GPS."),
    ("Re-anchor", "At a known corner the operator resets x, y and snaps the heading to the nearest 90°."),
    ("Place the mine", None),
    ("Merge and persist", "Hits within a merge radius are one mine; records survive a restart. The map is generated automatically."),
]
y = 1.85
for i, (h, b) in enumerate(steps):
    badge(s, 5.75, y + 0.02, i + 1, d=0.4, size=13)
    text(s, 6.3, y, 6.3, 0.4, h, size=16, color=NIGHT, bold=True)
    if b:
        text(s, 6.3, y + 0.38, 6.35, 0.75, b, size=13, color=MUTED)
        y += 1.15
    else:
        text(s, 6.3, y + 0.38, 6.35, 0.38, "Robot pose + detector offset (dx, dy), rotated by heading θ:",
             size=13, color=MUTED)
        rect(s, 6.3, y + 0.8, 6.35, 0.72, fill=TEAL_LIGHT)
        text(s, 6.45, y + 0.8, 6.1, 0.72, [
            [("xₘ = x + dx·cos θ − dy·sin θ", {"font": "Cambria"})],
            [("yₘ = y + dx·sin θ + dy·cos θ", {"font": "Cambria"})]],
            size=14, color=NIGHT, anchor=MSO_ANCHOR.MIDDLE)
        y += 1.65
notes(s, "19 × 19 map, columns A–S and rows 1–19, as the rulebook requires. Map "
         "frame: origin at the start corner, +X right, +Y up, heading 0 along +X.")

# 22 — Detection and classification
s = content_slide("Software", "Mine detection and classification")
box(s, 0.667, 1.9, 2.5, 1.0, [[("Metal detector", {"bold": True})], [("rising edge", {"size": 12, "color": MUTED})]],
    fill=WHITE, line=LINE, size=15)
box(s, 0.667, 3.35, 2.5, 1.0, [[("Camera → ONNX", {"bold": True})], [("fresh vision result?", {"size": 12, "color": MUTED})]],
    fill=WHITE, line=LINE, size=15)
box(s, 0.667, 4.8, 2.5, 1.0, [[("Light + siren", {"bold": True, "color": RED})],
                              [("on every metal hit", {"size": 12, "color": MUTED})]],
    fill=RGBColor(0xFB, 0xE3, 0xE3), line=None, size=15)
arrow(s, 1.92, 2.92, 1.92, 3.33, color=NIGHT)
arrow(s, 1.92, 4.37, 1.92, 4.78, color=NIGHT)
dec = [
    ("Metal + camera sees a box", "SURFACE", "Mapped, then the autonomous pickup sequence collects it", TEAL),
    ("Metal + camera sees nothing", "BURIED", "Mapped at its position; the arm is not used", SAND),
]
for i, (cond, cls, note, colr) in enumerate(dec):
    y = 1.9 + i * 1.95
    arrow(s, 3.2, 2.4 if i == 0 else 3.85, 3.6, y + 0.85, color=LINE, width=1.25)
    rect(s, 3.65, y, 9.0, 1.7, fill=WHITE, line=LINE)
    text(s, 3.85, y, 3.6, 1.7, cond, size=16, color=INK, anchor=MSO_ANCHOR.MIDDLE)
    box(s, 7.45, y + 0.5, 1.7, 0.7, cls, fill=colr, line=None, size=15, color=WHITE, bold=True)
    text(s, 9.35, y, 3.2, 1.7, note, size=14, color=MUTED, anchor=MSO_ANCHOR.MIDDLE)
rect(s, 3.65, 5.85, 9.0, 0.6, fill=TEAL_LIGHT)
text(s, 3.85, 5.85, 8.7, 0.6, "Vision alone never creates a mine: the metal detector must confirm it.",
     size=14, color=NIGHT, bold=True, anchor=MSO_ANCHOR.MIDDLE)
notes(s, "The single-class object detector runs in ONNX Runtime on the laptop.")

# 23 — Control station GUI
s = content_slide("Software", "Control station GUI")
if GUI.exists():
    picture(s, GUI, 0.667, 1.8, 7.75, 4.65, rounded=False)
else:
    rect(s, 0.667, 1.8, 7.75, 4.65, fill=RGBColor(0xFB, 0xF4, 0xE4), line=LINE)
    box(s, 0.8, 1.95, 2.55, 1.75, "Camera feed", fill=WHITE, line=RGBColor(0xD8, 0x1B, 0x60),
        color=RGBColor(0xD8, 0x1B, 0x60))
    box(s, 0.8, 3.8, 2.55, 2.0, "Controls", fill=WHITE, line=RGBColor(0x1E, 0x88, 0xE5),
        color=RGBColor(0x1E, 0x88, 0xE5))
    box(s, 3.5, 1.95, 4.8, 3.85, "Field map", fill=WHITE, line=TEAL, color=TEAL)
    box(s, 0.8, 5.9, 7.5, 0.45, "Warnings & messages", fill=WHITE, line=SAND, color=SAND)
callouts = [
    ("Camera feed", "Live RTSP video from the robot"),
    ("Field map", "Mines placed automatically as they are found"),
    ("Drive controls", "Arrow keys with smooth ramp; Space locks speed"),
    ("Pick up mine (J)", "Starts the autonomous pickup; disabled until the arm is upright"),
    ("Link health", "Stale data greyed out, never shown as zero"),
]
for i, (h, b) in enumerate(callouts):
    y = 1.8 + i * 0.95
    badge(s, 8.75, y + 0.05, i + 1, d=0.4, size=13)
    text(s, 9.3, y, 3.4, 0.4, h, size=15, color=NIGHT, bold=True)
    text(s, 9.3, y + 0.36, 3.4, 0.55, b, size=13, color=MUTED)
notes(s, "C++20 with Raylib for drawing and GStreamer for video. Layout and colours "
         "come from styles.toml, so they can be changed without recompiling.")

# 24 — Navigation plan
s = content_slide("Operation", "Arena navigation plan")
ax, ay, asz = 0.95, 1.95, 4.4
rect(s, ax, ay, asz, asz, fill=SAND_LIGHT, line=SAND, rounded=False)
lanes = 7
lane_w = asz / lanes
for k in range(lanes):
    x = ax + (k + 0.5) * lane_w
    y1, y2 = (ay + asz - 0.2, ay + 0.2) if k % 2 == 0 else (ay + 0.2, ay + asz - 0.2)
    arrow(s, x, y1, x, y2, color=TEAL, width=2)
    if k < lanes - 1:
        yy = y2
        arrow(s, x, yy, x + lane_w, yy, color=TEAL, width=2)
badge(s, ax - 0.2, ay + asz - 0.22, "S", d=0.42, fill=NIGHT)
for cx, cy in [(ax + asz - 0.22, ay - 0.2), (ax + asz - 0.22, ay + asz - 0.22), (ax - 0.2, ay - 0.2)]:
    badge(s, cx, cy, "R", d=0.42, fill=CYAN)
text(s, ax, ay + asz + 0.08, asz, 0.32, "S start corner · R pose reset at a corner",
     size=11, color=MUTED, align=PP_ALIGN.CENTER)
plan = [
    ("Start in a corner", "Pose reset to the known corner before the first lane."),
    ("Lawnmower sweep", "Parallel lanes, spacing = detector width minus overlap, so no strip of the field is missed."),
    ("Stop, signal, classify", "Light and siren fire on every metal hit; surface mines are picked up on the spot."),
    ("Re-anchor at every corner", "Each lane end near a corner resets drift."),
    ("Budget the 20 minutes", "Sweep time ≈ (19 m / lane spacing) × 19 m / speed, planned to cover the whole field."),
]
for i, (h, b) in enumerate(plan):
    y = 1.85 + i * 0.92
    badge(s, 5.9, y + 0.02, i + 1, d=0.4, size=13)
    text(s, 6.45, y, 6.2, 0.4, h, size=16, color=NIGHT, bold=True)
    text(s, 6.45, y + 0.37, 6.2, 0.55, b, size=13, color=MUTED)
notes(s, "Lane spacing and cruise speed come from field tests with the real "
         "detector.")

# 25 — Engineering quality
s = content_slide("Software", "Built to be reliable")
cards = [
    ("One contract", "A single versioned C++ header defines every message the GUI, bridge and vision process exchange."),
    ("Validate everything", "Wrong size, version, flag, range or non-finite float: the message is dropped."),
    ("Dockerized robot stack", "Host networking for DDS, devices found at boot, config bind-mounted so we can tune on site without rebuilding."),
    ("Tested", "Unit and integration tests cover drive ramping, pickup gating, message decoding and the ZeroMQ link."),
]
for i, (h, b) in enumerate(cards):
    col, row = i % 2, i // 2
    x, y = 0.667 + col * 6.1, 1.85 + row * 2.35
    rect(s, x, y, 5.9, 2.15, fill=WHITE, line=LINE)
    badge(s, x + 0.25, y + 0.25, i + 1, d=0.5, size=16, fill=TEAL if i % 3 == 0 else SAND)
    text(s, x + 0.95, y + 0.25, 4.7, 0.5, h, size=18, color=NIGHT, bold=True,
         anchor=MSO_ANCHOR.MIDDLE)
    text(s, x + 0.25, y + 0.95, 5.4, 1.1, b, size=14, color=MUTED)
notes(s, "Telemetry arrives at 20 Hz; any source silent for 500 ms is marked stale.")

# 26 — Team and next steps
s = content_slide("Team", "Team, budget and next steps")
rect(s, 0.667, 1.85, 4.0, 2.25, fill=TEAL_LIGHT)
text(s, 0.9, 1.95, 3.6, 2.05, [
    [("Watad Ejust", {"bold": True, "size": 26, "color": TEAL, "font": TITLE_FONT})],
    [("Academia category", {"size": 14, "color": MUTED})],
    [("Mechanical · electronics · software", {"size": 14, "color": INK})],
], spacing=4, anchor=MSO_ANCHOR.MIDDLE)
rect(s, 0.667, 4.3, 4.0, 2.15, fill=WHITE, line=LINE)
text(s, 0.9, 4.4, 3.6, 1.95, [
    [("Budget-minded build", {"bold": True, "size": 16, "color": NIGHT})],
    [("Wheelchair drive motors and a car wiper motor bought from the local market; "
      "an open-source software stack.", {"size": 14, "color": MUTED})],
], spacing=4)
picture(s, P8, 5.0, 1.85, 7.67, 2.6, focus=(0.5, 0.6))
rect(s, 5.0, 4.65, 7.67, 1.8, fill=WHITE, line=LINE)
text(s, 5.25, 4.75, 7.2, 0.45, "After the competition", size=16, color=NIGHT, bold=True)
bullets(s, 5.25, 5.25, 7.2, 1.15, [
    "Autonomous navigation: coverage planning on top of the existing localization",
    "Onboard vision on a Jetson; the GUI protocol stays the same",
], size=14, spacing=6)
notes(s, "Add team member names and the total spend if the jury asks.")

# ---------------------------------------------------------------- reorder
sld_ids = prs.slides._sldIdLst
items = list(sld_ids)
title_id, template_content_id, thanks_id = items[0], items[1], items[2]
new_ids = items[3:]
for el in items:
    sld_ids.remove(el)
for el in [title_id] + new_ids + [thanks_id]:
    sld_ids.append(el)
# Drop the unused template content slide.
prs.part.drop_rel(template_content_id.get(qn("r:id")))

prs.save(str(OUT))
print("saved", OUT, "slides:", len(prs.slides))
